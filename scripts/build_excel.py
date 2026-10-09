"""Build the TikTok review analysis Excel report with native charts.

Usage: python build_excel.py [reviews_json] [output_xlsx]
"""
import json
import re
import sys
import time
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else \
    Path(r"D:\code\codeMy\shop-qc\probe\us_cdp\all_reviews.json")
OUT = Path(r"D:\code\codeMy\shop-qc\report")
OUT.mkdir(parents=True, exist_ok=True)
XLSX = Path(sys.argv[2]) if len(sys.argv) > 2 else \
    OUT / "TikTok评论分析报告_单肩包1729483763488494309.xlsx"

# palette (star polarity ramp + single hue for neutral series)
STAR_COLORS = {"5": "2E7D32", "4": "7CB342", "3": "F9A825", "2": "EF6C00", "1": "C62828"}
ACCENT = "FE2C55"      # TikTok pink for neutral single-series bars
INK = "1F2328"
MUTED = "6B7280"
HEAD_FILL = PatternFill("solid", fgColor="FFF0F3")
THIN = Side(style="thin", color="E5E7EB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

data = json.loads(SRC.read_text(encoding="utf-8"))
revs = data["reviews"]


def ts_month(ms):
    try:
        return time.strftime("%Y-%m", time.localtime(int(ms) / 1000))
    except Exception:
        return None


def norm_sku(s):
    if not s:
        return "未标注"
    key = re.sub(r"\s+", " ", s).strip().lower()
    fixes = {"coffee brown, one-size": "Coffee Brown, One-size",
             "coffee brown, please refer to the pictures for dimensions":
             "Coffee Brown, 尺寸见图",
             "please refer to the pictures for dimensions": "尺寸见图"}
    return fixes.get(key, s.strip().title() if len(s) < 40 else s.strip())


n = len(revs)
if n == 0:
    print("FATAL: 0 reviews harvested - nothing to report")
    sys.exit(2)
avg_overall = data.get("overall_score")
star_cnt = Counter(r["review_rating"] for r in revs)
neg = [r for r in revs if r["review_rating"] <= 2]
pos5 = [r for r in revs if r["review_rating"] == 5]
with_img = [r for r in revs if r.get("review_images")]
verified = sum(1 for r in revs if r.get("is_verified_purchase"))
incent = sum(1 for r in revs if r.get("is_incentivized_review"))

# monthly trend
months = {}
for r in revs:
    m = ts_month(r.get("review_time"))
    if not m:
        continue
    months.setdefault(m, []).append(r["review_rating"])
mkeys = sorted(months)

# SKU
sku_cnt = Counter(norm_sku(r.get("sku_specification")) for r in revs)
sku_rate = {}
for r in revs:
    k = norm_sku(r.get("sku_specification"))
    sku_rate.setdefault(k, []).append(r["review_rating"])

# negative keywords
STOP = set("""the a an and or but i im i'm it its it's is was were be been being to of in on for
with my me you your they them he she we us this that these those so very just not no yes as at
by from have has had do does did will would can could should about if then than there here what
when who how why all one two get got really too more most much many out up down over again only
bought purchase purchased ordered order using use used love loved like loves nice good great
wouldnt wouldnt didnt dont don't doesnt doesn't isnt isn't wasnt wasn't cant can't ive i've
""".split())
words = Counter()
for r in neg:
    for w in re.findall(r"[a-z']{3,}", (r.get("review_text") or "").lower()):
        if w not in STOP:
            words[w] += 1
neg_top = words.most_common(15)

wb = Workbook()

# ---------- helpers ----------
def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True, color=INK)
        cell.fill = HEAD_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")


def kpi(ws, row, col, label, value, sub=""):
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
    ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)
    ws.merge_cells(start_row=row + 2, start_column=col, end_row=row + 2, end_column=col + 1)
    ws.cell(row=row, column=col, value=label).font = Font(size=10, color=MUTED)
    v = ws.cell(row=row + 1, column=col, value=value)
    v.font = Font(size=22, bold=True, color=INK)
    s = ws.cell(row=row + 2, column=col, value=sub)
    s.font = Font(size=9, color=MUTED)


# ---------- Sheet 1: 总览 ----------
ws = wb.active
ws.title = "总览"
ws.sheet_view.showGridLines = False
ws["A1"] = "TikTok Shop 商品评论分析报告"
ws["A1"].font = Font(size=15, bold=True, color=INK)
ws["A2"] = (f"商品: {(revs[0]['product_name'][:60] if revs else 'N/A')}  |  商品ID: {data['product_id']}  |  "
            f"抓取时间: {data['harvested_at']}  |  样本: {n}/{data['total_reviews_claimed']} 条")
ws["A2"].font = Font(size=9, color=MUTED)

kpi(ws, 4, 1, "评论总数", n, f"平台显示 {data['total_reviews_claimed']}")
kpi(ws, 4, 3, "综合评分", f"{avg_overall}★", "满分 5 星")
kpi(ws, 4, 5, "5星占比", f"{star_cnt[5]/n:.0%}", f"{star_cnt[5]} 条")
kpi(ws, 4, 7, "差评(1-2星)", len(neg), f"差评率 {len(neg)/n:.1%}")
kpi(ws, 4, 9, "带图评论", len(with_img), f"占比 {len(with_img)/n:.0%}")
kpi(ws, 4, 11, "真实购买", f"{verified/n:.0%}", f"{verified} 条")
kpi(ws, 4, 13, "激励评论", incent, f"占比 {incent/n:.0%}")

# star table + pie
ws["A12"] = "星级分布"
ws["A12"].font = Font(size=12, bold=True)
ws.cell(row=13, column=1, value="星级").font = Font(bold=True)
ws.cell(row=13, column=2, value="数量").font = Font(bold=True)
ws.cell(row=13, column=3, value="占比").font = Font(bold=True)
for i, s in enumerate([5, 4, 3, 2, 1]):
    ws.cell(row=14 + i, column=1, value=f"{s}星")
    ws.cell(row=14 + i, column=2, value=star_cnt[s])
    ws.cell(row=14 + i, column=3, value=f"{star_cnt[s]/n:.1%}")
    for c in range(1, 4):
        ws.cell(row=14 + i, column=c).border = BORDER

pie = PieChart()
pie.title = "星级占比"
pie.height, pie.width = 8.2, 13
pie.add_data(Reference(ws, min_col=2, min_row=13, max_row=19), titles_from_data=True)
pie.set_categories(Reference(ws, min_col=1, min_row=14, max_row=19))
series = pie.series[0]
for i, s in enumerate(["5", "4", "3", "2", "1"]):
    dp = DataPoint(idx=i)
    dp.graphicalProperties.solidFill = STAR_COLORS[s]
    series.data_points.append(dp)
pie.dataLabels = DataLabelList()
pie.dataLabels.showPercent = True
pie.legend.position = "r"
ws.add_chart(pie, "E12")

# star bar (counts)
bar = BarChart()
bar.type, bar.style = "col", 10
bar.title = "各星级评论数"
bar.height, bar.width = 8.2, 13
bar.add_data(Reference(ws, min_col=2, min_row=13, max_row=19), titles_from_data=True)
bar.set_categories(Reference(ws, min_col=1, min_row=14, max_row=19))
bar.series[0].graphicalProperties.solidFill = ACCENT
bar.dataLabels = DataLabelList()
bar.dataLabels.showVal = True
bar.legend = None
bar.y_axis.majorGridlines = None
ws.add_chart(bar, "M12")

for col, w in {"A": 10, "B": 10, "C": 10, "D": 3}.items():
    ws.column_dimensions[col].width = w

# ---------- Sheet 2: 月度趋势 ----------
ws2 = wb.create_sheet("月度趋势")
ws2.sheet_view.showGridLines = False
ws2["A1"] = "月度评论量与平均星级"
ws2["A1"].font = Font(size=12, bold=True)
ws2.append([])
ws2.append(["月份", "评论量", "平均星级"])
hdr_row = 3
style_header(ws2, hdr_row, 3)
for i, m in enumerate(mkeys):
    vals = months[m]
    ws2.append([m, len(vals), round(sum(vals) / len(vals), 2)])
last = hdr_row + len(mkeys)
for row in ws2.iter_rows(min_row=hdr_row + 1, max_row=last, max_col=3):
    for c in row:
        c.border = BORDER

vbar = BarChart()
vbar.type, vbar.style = "col", 10
vbar.title = "每月评论量"
vbar.height, vbar.width = 9, 20
vbar.add_data(Reference(ws2, min_col=2, min_row=hdr_row, max_row=last), titles_from_data=True)
vbar.set_categories(Reference(ws2, min_col=1, min_row=hdr_row + 1, max_row=last))
vbar.series[0].graphicalProperties.solidFill = ACCENT
vbar.dataLabels = DataLabelList()
vbar.dataLabels.showVal = True
vbar.legend = None
vbar.y_axis.majorGridlines = None
ws2.add_chart(vbar, "E3")

line = LineChart()
line.title = "每月平均星级"
line.height, line.width = 9, 20
line.add_data(Reference(ws2, min_col=3, min_row=hdr_row, max_row=last), titles_from_data=True)
line.set_categories(Reference(ws2, min_col=1, min_row=hdr_row + 1, max_row=last))
s = line.series[0]
s.graphicalProperties.line.solidFill = "2E7D32"
s.graphicalProperties.line.width = 28000
s.marker.symbol = "circle"
s.marker.size = 7
line.y_axis.scaling.min = 3.5
line.y_axis.scaling.max = 5
line.dataLabels = DataLabelList()
line.dataLabels.showVal = True
line.legend = None
ws2.add_chart(line, "E22")
ws2.column_dimensions["A"].width = 12

# ---------- Sheet 3: SKU分析 ----------
ws3 = wb.create_sheet("SKU分析")
ws3.sheet_view.showGridLines = False
ws3["A1"] = "SKU 规格分布与口碑"
ws3["A1"].font = Font(size=12, bold=True)
ws3.append([])
ws3.append(["SKU 规格", "评论量", "占比", "平均星级"])
style_header(ws3, 3, 4)
top_skus = sku_cnt.most_common()
for k, v in top_skus:
    ws3.append([k, v, f"{v/n:.1%}", round(sum(sku_rate[k]) / len(sku_rate[k]), 2)])
last3 = 3 + len(top_skus)
for row in ws3.iter_rows(min_row=4, max_row=last3, max_col=4):
    for c in row:
        c.border = BORDER

skubar = BarChart()
skubar.type, skubar.style = "bar", 10
skubar.title = "SKU 评论量(前10)"
skubar.height, skubar.width = 10, 16
top10 = min(10, len(top_skus))
skubar.add_data(Reference(ws3, min_col=2, min_row=3, max_row=3 + top10), titles_from_data=True)
skubar.set_categories(Reference(ws3, min_col=1, min_row=4, max_row=3 + top10))
skubar.series[0].graphicalProperties.solidFill = ACCENT
skubar.dataLabels = DataLabelList()
skubar.dataLabels.showVal = True
skubar.legend = None
skubar.y_axis.majorGridlines = None
ws3.add_chart(skubar, "F3")

skurate = BarChart()
skurate.type, skurate.style = "bar", 10
skurate.title = "各SKU平均星级(前10)"
skurate.height, skurate.width = 10, 16
skurate.add_data(Reference(ws3, min_col=4, min_row=3, max_row=3 + top10), titles_from_data=True)
skurate.set_categories(Reference(ws3, min_col=1, min_row=4, max_row=3 + top10))
skurate.series[0].graphicalProperties.solidFill = "2E7D32"
skurate.dataLabels = DataLabelList()
skurate.dataLabels.showVal = True
skurate.legend = None
skurate.y_axis.majorGridlines = None
skurate.x_axis.scaling.min = 3
skurate.x_axis.scaling.max = 5
ws3.add_chart(skurate, "F24")
ws3.column_dimensions["A"].width = 42

# ---------- Sheet 4: 差评分析 ----------
ws4 = wb.create_sheet("差评分析")
ws4.sheet_view.showGridLines = False
ws4["A1"] = f"差评分析（1-2星，共 {len(neg)} 条，占 {len(neg)/n:.1%}）"
ws4["A1"].font = Font(size=12, bold=True)
ws4["A2"] = "差评高频词（英文，已去除停用词）"
ws4["A2"].font = Font(size=10, color=MUTED)
ws4.append([])
ws4.append(["关键词", "出现次数"])
kw_hdr = 4
style_header(ws4, kw_hdr, 2)
for w, c in neg_top:
    ws4.append([w, c])
for row in ws4.iter_rows(min_row=kw_hdr + 1, max_row=kw_hdr + len(neg_top), max_col=2):
    for c in row:
        c.border = BORDER

kwbar = BarChart()
kwbar.type, kwbar.style = "bar", 10
kwbar.title = "差评高频词 Top15"
kwbar.height, kwbar.width = 12, 15
kwbar.add_data(Reference(ws4, min_col=2, min_row=kw_hdr, max_row=kw_hdr + len(neg_top)),
               titles_from_data=True)
kwbar.set_categories(Reference(ws4, min_col=1, min_row=kw_hdr + 1, max_row=kw_hdr + len(neg_top)))
kwbar.series[0].graphicalProperties.solidFill = "C62828"
kwbar.dataLabels = DataLabelList()
kwbar.dataLabels.showVal = True
kwbar.legend = None
kwbar.y_axis.majorGridlines = None
ws4.add_chart(kwbar, "D4")

start = kw_hdr + len(neg_top) + 3
ws4.cell(row=start - 1, column=1, value="差评明细").font = Font(size=11, bold=True)
ws4.append([])
for i, h in enumerate(["星级", "日期", "SKU", "评论内容", "图片链接"], start=1):
    ws4.cell(row=start, column=i, value=h)
style_header(ws4, start, 5)
for j, r in enumerate(sorted(neg, key=lambda x: x["review_time"], reverse=True), start=1):
    d = time.strftime("%Y-%m-%d", time.localtime(int(r["review_time"]) / 1000))
    ws4.cell(row=start + j, column=1, value=r["review_rating"])
    ws4.cell(row=start + j, column=2, value=d)
    ws4.cell(row=start + j, column=3, value=norm_sku(r.get("sku_specification")))
    ws4.cell(row=start + j, column=4, value=r.get("review_text"))
    ws4.cell(row=start + j, column=5, value=" | ".join(r.get("review_images") or []))
ws4.column_dimensions["A"].width = 10
ws4.column_dimensions["B"].width = 12
ws4.column_dimensions["C"].width = 22
ws4.column_dimensions["D"].width = 70
ws4.column_dimensions["E"].width = 40

# ---------- Sheet 5: 评论明细 ----------
ws5 = wb.create_sheet("评论明细")
headers = ["评论ID", "星级", "日期", "用户名", "国家", "真实购买", "激励评论",
           "SKU规格", "评论内容", "图片数", "图片链接"]
ws5.append(headers)
style_header(ws5, 1, len(headers))
for r in sorted(revs, key=lambda x: int(x["review_time"]), reverse=True):
    imgs = r.get("review_images") or []
    ws5.append([
        r["review_id"], r["review_rating"],
        time.strftime("%Y-%m-%d", time.localtime(int(r["review_time"]) / 1000)),
        r.get("reviewer_name"), r.get("review_country"),
        "是" if r.get("is_verified_purchase") else "否",
        "是" if r.get("is_incentivized_review") else "否",
        norm_sku(r.get("sku_specification")), r.get("review_text"),
        len(imgs), " | ".join(imgs[:3])])
widths = [20, 7, 12, 14, 8, 9, 9, 24, 60, 7, 40]
for i, w in enumerate(widths, start=1):
    ws5.column_dimensions[get_column_letter(i)].width = w
ws5.freeze_panes = "A2"
ws5.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{n + 1}"

wb.save(XLSX)
print("saved:", XLSX)
print(f"sheets: {[s.title for s in wb.worksheets]}")
print(f"rows detail: {n}, months: {len(mkeys)}, skus: {len(top_skus)}, neg: {len(neg)}")
